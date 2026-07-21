#!/usr/bin/env python3
"""Build the illustrated FakeBuster setup and operations guide."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "output" / "playwright" / "fakebuster"
OUTPUT = ROOT / "docs" / "FakeBuster-workflow-setup.docx"

INK = "16302A"
GREEN = "0E7459"
PALE = "EAF4F0"
GOLD = "D69E2E"
MUTED = "5F6F69"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Liberation Sans"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("TubeFactory - FakeBuster setup  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_bullets(document: Document, values: list[str], *, style: str = "List Bullet") -> None:
    for value in values:
        paragraph = document.add_paragraph(style=style)
        paragraph.add_run(value)


def add_numbered(document: Document, values: list[str]) -> None:
    for index, value in enumerate(values, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.65)
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(value)


def add_note(document: Document, title: str, body: str, *, accent: str = GREEN) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, PALE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(title + "  ")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(accent)
    body_run = paragraph.add_run(body)
    body_run.font.color.rgb = RGBColor.from_string(INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(
    document: Document,
    filename: str,
    caption: str,
    *,
    width: float,
) -> None:
    path = SCREENSHOTS / filename
    if not path.exists():
        raise FileNotFoundError(f"Missing required screenshot: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(8)
    caption_run.font.color.rgb = RGBColor.from_string(MUTED)


def add_table(document: Document, rows: list[tuple[str, str]], headers: tuple[str, str]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, value in enumerate(headers):
        shade(table.rows[0].cells[index], GREEN)
        set_cell_text(table.rows[0].cells[index], value, bold=True, color=WHITE)
    for left, right in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], left, bold=True)
        set_cell_text(cells[1], right)
        shade(cells[0], "F4F7F5")
    document.add_paragraph()


def page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 30, INK),
        ("Subtitle", 13, MUTED),
        ("Heading 1", 20, INK),
        ("Heading 2", 13, GREEN),
        ("Heading 3", 10.5, INK),
    ):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(5)

    styles["List Bullet"].font.name = "Liberation Sans"
    styles["List Number"].font.name = "Liberation Sans"

    header = section.header.paragraphs[0]
    header.text = "FAKEBUSTER  /  VERIFIED OPERATIONS GUIDE"
    header.runs[0].font.name = "Liberation Sans"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    core = document.core_properties
    core.title = "FakeBuster workflow setup"
    core.subject = "TubeFactory scheduled German misinformation discovery and review workflow"
    core.author = "TubeFactory operations"
    core.keywords = "FakeBuster, TubeFactory, misinformation, evidence, workflow"


def build() -> None:
    document = Document()
    configure_document(document)

    # Cover
    document.add_paragraph("TUBEFACTORY", style="Subtitle")
    title = document.add_paragraph(style="Title")
    title.add_run("FakeBuster\nworkflow setup")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Automated German news discovery, evidence-first triage and human approval")
    document.add_paragraph()
    band = document.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    shade(band.cell(0, 0), GREEN)
    band_text = band.cell(0, 0).paragraphs[0]
    band_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = band_text.add_run("VALIDATED 21 JULY 2026  |  DEVELOPMENT BRANCH: dev")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(WHITE)
    document.add_paragraph()
    document.add_paragraph(
        "This guide records the exact setup that produced a working scheduled FakeBuster workflow in TubeFactory. "
        "It includes the operator-facing configuration, browser evidence, acceptance results, safeguards and the defects corrected during commissioning."
    )
    add_note(
        document,
        "Editorial boundary",
        "Automation may label an item only as potential misinformation. The stronger conclusion demonstrably false or materially misleading requires an evidence dossier and explicit human approval.",
        accent=GOLD,
    )
    document.add_paragraph()
    add_table(
        document,
        [
            ("Control plane", "http://evidence-studio.local:8090"),
            ("Workspace", "FakeBuster (enabled, language de)"),
            ("Administrator account", "gate-admin; password is intentionally not stored in this document or repository"),
            ("Schedule", "Every 6 hours, Europe/Berlin"),
            ("Freshness window", "7 days"),
            ("Runtime", "Docker Compose project youtuber; Temporal scheduled workflow"),
        ],
        ("Item", "Verified value"),
    )

    page_break(document)
    document.add_heading("1. What the workflow does", level=1)
    document.add_paragraph(
        "FakeBuster continuously searches for recently published German-language reporting that may contain a materially important, verifiable factual claim. Discovery emphasizes effects on democracy, public safety, peace and security, health, the economy, the environment and critical infrastructure."
    )
    document.add_heading("Decision sequence", level=2)
    add_numbered(
        document,
        [
            "Discover recent German-language results using the subject's search plan and falsification branch.",
            "Apply exclusions for satire, opinion, commentary, columns and predictions.",
            "Normalize URLs, remove exact result duplicates, remember processed URLs and suppress substantially identical title/excerpt claim clusters across runs.",
            "Create ranked pending opportunities labelled potential misinformation.",
            "Require an editor to inspect the exact disputed claim, original source, dates, syndication and uncertainties.",
            "Acquire immutable source snapshots and build a claim-by-claim evidence dossier.",
            "Permit the final stronger finding only when at least two independent sources, including at least one primary source, substantively contradict the central claim and a human reviewer approves it.",
        ],
    )
    add_note(
        document,
        "Not a controversy detector",
        "Bias, exaggeration, poor wording, minority opinion, marked satire, commentary, predictions and immaterial errors are not enough. A central factual claim must be demonstrably false, fabricated, materially misleading, out of context or missing essential context.",
    )
    document.add_heading("What counts as already processed", level=2)
    add_bullets(
        document,
        [
            "The same canonical URL, after tracking parameters and URL variants are normalized.",
            "The same normalized title and search excerpt fingerprint.",
            "A different URL with highly similar title/excerpt tokens representing the same central claim cluster.",
            "After acquisition, identical full-content hashes or high-similarity content chunks, which are also used to identify dependent or syndicated sources.",
            "A materially different update remains eligible when its title and excerpt contain genuinely new claims or evidence.",
        ],
    )

    page_break(document)
    document.add_heading("2. Create and enable the channel", level=1)
    document.add_paragraph(
        "Open Channels & subjects. Create the channel in a disabled state first so its identity can be reviewed before any work is attached to it. Use the values below."
    )
    add_table(
        document,
        [
            ("Channel name", "FakeBuster"),
            ("Slug", "fakebuster"),
            ("Languages", "de"),
            ("Description", "German evidence-based fact-check channel; potential misinformation discovery with human approval"),
        ],
        ("Field", "Value"),
    )
    add_figure(document, "02-channel-form.png", "Figure 1 - Channel form completed before creation.", width=4.2)
    document.add_paragraph(
        "Select Create disabled channel. Confirm that FakeBuster appears as Disabled, then select Enable. This two-step sequence prevents accidental scheduling against an unreviewed channel identity."
    )
    add_figure(document, "03-channel-created-disabled.png", "Figure 2 - The new channel exists but is safely disabled before review.", width=3.0)
    add_note(
        document,
        "Activation checkpoint",
        "The screenshot records the deliberate disabled state. Enabling produces version 2 and makes FakeBuster available as the active channel workspace; Disable remains available as the safe rollback.",
    )

    page_break(document)
    document.add_heading("3. Create the monitoring subject", level=1)
    document.add_paragraph(
        "With FakeBuster selected as the channel workspace, create a disabled subject using these values. Keep the research goal explicit: it is the durable editorial instruction used to generate the search plan."
    )
    add_table(
        document,
        [
            ("Profile name", "FakeBuster German misinformation monitor"),
            ("Language / risk", "de / high"),
            ("Operating profile", "Assisted"),
            ("Recent article window", "7 days"),
            ("Timezone", "Europe/Berlin"),
            ("Temporal cron", "0 */6 * * *"),
            ("Independent / primary minimum", "2 / 1"),
            ("Excluded terms", "Satire, Meinung, Kommentar, Kolumne, Prognose"),
            ("Sensitive categories", "misinformation fact-checking, politics, breaking news, health, finance, identifiable accusation, legal"),
        ],
        ("Field", "Configured value"),
    )
    document.add_heading("Seed-query themes", level=2)
    add_bullets(
        document,
        [
            "German-language current news plus false report / fact-check terms.",
            "Claims contradicted by an official source.",
            "Disinformation affecting democracy, safety, health, the economy or the environment.",
            "Allegedly false, misleading or out-of-context reporting.",
            "Corrections, retractions and clarifications in German reporting.",
        ],
    )
    document.add_heading("Configured form evidence", level=2)
    add_figure(document, "05-subject-form.png", "Figure 3 - Complete subject form, including exclusions, recency and schedule.", width=2.25)

    page_break(document)
    document.add_heading("4. Validate the search plan before enabling it", level=1)
    document.add_paragraph(
        "Select Test plan while the subject remains disabled. Review every broad-discovery query, the primary-evidence query and the falsification branch. The generated plan must be German-language, retain the exclusions and actively search for correction or counterevidence."
    )
    add_figure(document, "07-validated-search-plan.png", "Figure 4 - Validated German search plan with primary-evidence and falsification branches.", width=6.35)
    add_note(
        document,
        "Go/no-go check",
        "Do not enable or schedule a subject whose plan omits counterevidence, uses the wrong language, or silently drops the satire/opinion exclusions.",
        accent=GOLD,
    )

    page_break(document)
    document.add_heading("5. Enable and schedule automatic discovery", level=1)
    add_numbered(
        document,
        [
            "Select Enable on the validated subject.",
            "Select Apply schedule. The status must read schedule active and show the next execution time.",
            "Leave the channel and subject enabled. Temporal now launches discovery every six hours in Europe/Berlin.",
            "Use Workflow activity to inspect completed, running or attention-needed runs without remembering workflow IDs.",
        ],
    )
    add_figure(document, "09-schedule-active.png", "Figure 5 - Subject enabled and its Temporal schedule active.", width=4.4)
    document.add_heading("Automated run behavior", level=2)
    add_bullets(
        document,
        [
            "SearXNG receives a time-range filter derived from the 7-day freshness policy.",
            "Search responses are read to end-of-stream, with bounded response size and Temporal retry support.",
            "A subject-level database lock prevents concurrent runs from creating the same opportunity.",
            "Canonical URLs and prior claim signatures are checked before new opportunities are written.",
            "Publisher and publication timestamps are retained for later provenance review.",
            "Every scheduled run is recorded in the same durable workflow activity and retry surface as manual live discovery.",
        ],
    )
    document.add_heading("Workflow activity evidence", level=2)
    add_figure(document, "13-final-scheduled-run-completed.png", "Figure 6 - Final scheduled acceptance run completed and visible in Workflow activity.", width=6.25)

    page_break(document)
    document.add_heading("6. Triage candidates without overclaiming", level=1)
    document.add_paragraph(
        "Discovery results enter the opportunity board as pending. The system deliberately shows the neutral discovery label potential misinformation. This is a lead for investigation, not a finding that the publisher or claim is false."
    )
    add_figure(document, "11-potential-misinformation-candidate.png", "Figure 7 - A pending candidate uses the required neutral discovery label.", width=5.2)
    document.add_heading("Required operator review", level=2)
    add_numbered(
        document,
        [
            "Write the exact disputed factual claim. Do not substitute a general impression of the article.",
            "Record what supports and contradicts the claim, including the original source, publication date and event date.",
            "Check whether multiple reports are independent or merely syndicated copies.",
            "Record uncertainties and credible alternative explanations.",
            "Shortlist only when the claim is materially important and plausibly testable with reliable evidence.",
            "Acquire sources, inspect immutable snapshots and build the dossier.",
            "Approve the stronger conclusion only after the evidence threshold and human-review gate are satisfied.",
        ],
    )
    add_table(
        document,
        [
            ("Minimum evidence", "At least 2 independent contradicting sources"),
            ("Primary evidence", "At least 1 primary source"),
            ("Contradiction", "Substantive and directed at the central factual claim"),
            ("Final gate", "Explicit human approval; no automatic final accusation"),
            ("Ranked record", "Title, publisher, URL, date, relevance, disputed claim, evidence summary, contradiction status, independence, confidence and action"),
        ],
        ("Gate", "Requirement"),
    )

    page_break(document)
    document.add_heading("7. Use the searchable task help", level=1)
    document.add_paragraph(
        "Select Help in the left navigation for longer procedures. The panel is keyboard-focusable, scoped to the signed-in role and searchable by goal, stage, title, summary or step text. Expand a result to see the ordered procedure, then use its action button to open the relevant stage."
    )
    add_figure(document, "14-searchable-task-help.png", "Figure 8 - Searching for channel narrows the procedures and reveals the five-step channel/subject guide.", width=4.65)
    add_note(
        document,
        "Two levels of guidance",
        "Hover or focus individual controls for concise field-level instructions. Use the Help panel for longer task-based procedures that cross several controls or stages.",
    )

    page_break(document)
    document.add_heading("8. Issues found and corrected", level=1)
    document.add_paragraph(
        "The following defects were encountered during commissioning. Each fix was deployed and rechecked against the live workflow."
    )
    add_table(
        document,
        [
            ("Channel could be created but not enabled", "Added version-checked Enable/Disable controls for channel profiles."),
            ("Freshness silently used 30 days", "Added the recent-article field and propagated lookback_days through workflow activities to the SearXNG time_range."),
            ("Excluded terms were not captured", "Added the form field and persisted negative keywords into every generated search query."),
            ("Repeated URLs returned on later runs", "Added subject-scoped cross-run canonical URL checks under an advisory transaction lock."),
            ("Same claim could reappear under another URL", "Added normalized title/excerpt fingerprints and a conservative high-similarity processed-claim check; materially different updates remain eligible."),
            ("Publisher and publication time were dropped", "Preserved publisher/domain metadata and normalized SearXNG publication timestamps."),
            ("SearXNG JSON was occasionally truncated", "Changed the client from a single bounded read to bounded chunked reading until end-of-stream; Temporal retries remain in place."),
            ("Scheduled runs were missing from Workflow activity", "Inserted scheduled executions into the durable live-discovery control surface with workflow ID and correlation lineage."),
            ("A failed run left its Temporal schedule paused", "Schedule reconciliation now explicitly unpauses enabled schedules and pauses disabled schedules."),
            ("Candidates could look like final accusations", "Added the visible potential misinformation classification and retained human approval for any stronger label."),
        ],
        ("Observed problem", "Correction"),
    )

    document.add_heading("Verification evidence", level=2)
    add_table(
        document,
        [
            ("API test suite", "89 passed"),
            ("Research-worker test suite", "36 passed"),
            ("Web", "TypeScript and production Next.js build passed; browser exercised with Playwright"),
            ("Final workflow ID", "scheduled-subject-discovery-99a8655d-4ac1-4f34-9294-40419052ea03-2026-07-21T14:35:44Z"),
            ("Final state", "Completed; no API or research-worker errors in the validation window"),
            ("Final counts", "36 raw; 35 unique; 1 within-run duplicate; 23 previously processed; 12 new candidates; 11 source domains"),
            ("Schedule", "Active; every 6 hours in Europe/Berlin"),
        ],
        ("Check", "Observed result"),
    )

    page_break(document)
    document.add_heading("9. Routine operator checklist", level=1)
    document.add_heading("Each shift", level=2)
    add_bullets(
        document,
        [
            "Open the FakeBuster workspace and check Workflow activity for Needs attention.",
            "Confirm the subject card still says schedule active.",
            "Review new pending opportunities, beginning with material impact and a precisely testable central claim.",
            "Reject satire, opinion, predictions, trivial errors and unsupported controversy leads early.",
            "Never treat source count as independence; inspect syndication and shared text.",
        ],
    )
    document.add_heading("When a run fails", level=2)
    add_numbered(
        document,
        [
            "Open Workflow activity and filter Needs attention.",
            "Open the failed stage and read the last durable status and workflow log.",
            "Correct the underlying network, source or configuration issue.",
            "Use Retry as new run when it is offered so lineage and the original failure remain auditable.",
            "Return to Channels & subjects and select Apply schedule if the schedule was failure-paused; confirm schedule active afterward.",
        ],
    )
    document.add_heading("Security and audit notes", level=2)
    add_bullets(
        document,
        [
            "Do not store the gate-admin password in documentation, source control or screenshots.",
            "All shortlist, rejection, dossier and final-media decisions require a reason and remain auditable.",
            "Source acquisition uses public-address validation, redirect checks, robots policy, immutable hashes and bounded content handling.",
            "The discovery label is intentionally cautious; accusations about identifiable people or organizations receive human review.",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
