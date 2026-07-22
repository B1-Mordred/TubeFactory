#!/usr/bin/env python3
"""Build the illustrated FaktischSimpel automation workflow guide."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "output" / "playwright" / "faktischsimpel"
BLUEPRINT = ROOT / "config" / "channel-workflows" / "faktischsimpel.json"
OUTPUT = ROOT / "docs" / "FaktischSimpel-Automatisierungsworkflow.docx"

INK = "142C35"
TEAL = "0D6F70"
BLUE = "2D5B87"
PALE = "EAF4F4"
GOLD = "B7791F"
MUTED = "5E6E74"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Liberation Sans"
    run.font.size = Pt(8.8)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("TubeFactory – FaktischSimpel  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def add_table(document: Document, rows: list[tuple[str, str]], headers: tuple[str, str]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, value in enumerate(headers):
        shade(table.rows[0].cells[index], TEAL)
        set_cell_text(table.rows[0].cells[index], value, bold=True, color=WHITE)
    for left, right in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], left, bold=True)
        set_cell_text(cells[1], right)
        shade(cells[0], "F3F7F7")
    document.add_paragraph()


def add_note(document: Document, title: str, body: str, *, accent: str = TEAL) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, PALE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(title + "  ")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(accent)
    paragraph.add_run(body)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(document: Document, values: list[str]) -> None:
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def add_numbered(document: Document, values: list[str]) -> None:
    for index, value in enumerate(values, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.65)
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(value)


def add_figure(document: Document, filename: str, caption: str, *, width: float) -> None:
    path = SCREENSHOTS / filename
    if not path.exists():
        raise FileNotFoundError(f"Missing required screenshot: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(8)
    caption_run.font.color.rgb = RGBColor.from_string(MUTED)


def add_prompt(document: Document, title: str, prompt: str) -> None:
    document.add_heading(title, level=2)
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "F3F7F8")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(prompt)
    run.font.name = "Liberation Mono"
    run.font.size = Pt(7.8)
    run.font.color.rgb = RGBColor.from_string(INK)
    document.add_paragraph()


def page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 29, INK),
        ("Subtitle", 12.5, MUTED),
        ("Heading 1", 19, INK),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 10.5, BLUE),
    ):
        style = document.styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(5)

    document.styles["List Bullet"].font.name = "Liberation Sans"
    header = section.header.paragraphs[0]
    header.text = "FAKTISCHSIMPEL  /  AUTOMATISIERTER ERKLÄRVIDEO-WORKFLOW"
    header.runs[0].font.name = "Liberation Sans"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section.footer.paragraphs[0])

    core = document.core_properties
    core.title = "FaktischSimpel Automatisierungsworkflow"
    core.subject = "TubeFactory channel-scoped evidence-first explainer workflow"
    core.author = "TubeFactory operations"
    core.keywords = "FaktischSimpel, TubeFactory, Erklärvideo, Evidenz, AI prompts, Workflow"


def build() -> None:
    blueprint = json.loads(BLUEPRINT.read_text(encoding="utf-8"))
    workflow = blueprint["channel"]["editorial_rules"]["automation_workflow"]
    subject = blueprint["subject"]
    prompts = workflow["prompts"]

    document = Document()
    configure_document(document)

    document.add_paragraph("TUBEFACTORY", style="Subtitle")
    title = document.add_paragraph(style="Title")
    title.add_run("FaktischSimpel\nAutomatisierungsworkflow")
    document.add_paragraph(
        "Von der Themenbeobachtung über beleggebundene AI-Skripte und unabhängige Prüfung bis zum privaten Upload",
        style="Subtitle",
    )
    document.add_paragraph()
    band = document.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    shade(band.cell(0, 0), TEAL)
    band_text = band.cell(0, 0).paragraphs[0]
    band_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = band_text.add_run("LIVE VALIDIERT: 22. JULI 2026  |  BRANCH: dev  |  WORKFLOW: v2")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(WHITE)
    document.add_paragraph()
    document.add_paragraph(
        "Dieses Dokument hält die implementierte und live konfigurierte FaktischSimpel-Pipeline fest. "
        "Sie automatisiert wiederkehrende Recherche- und Produktionsschritte, ohne redaktionelle Freigaben, "
        "Quellenbindung oder die private Standardveröffentlichung aufzuweichen."
    )
    add_note(
        document,
        "Kurzantwort",
        "Ja. TubeFactory kann geeignete Themen regelmäßig entdecken, Quellen verarbeiten, ein kanaltypisches deutsches Erklärskript schreiben, dieses unabhängig prüfen, ein Storyboard erzeugen und Medien vorbereiten. Redaktionelle Entscheidungen und die öffentliche Veröffentlichung bleiben bewusst menschliche Freigaben.",
    )
    add_table(
        document,
        [
            ("Kanal", "FaktischSimpel, aktiv, Version 3, Sprache de"),
            ("Workflow", "faktischsimpel.explainer, 13 Stufen, aktiv"),
            ("Themenradar", "FaktischSimpel Themenradar, aktiv, Version 3"),
            ("Zeitplan", "Montags 07:00 Uhr, Europe/Berlin"),
            ("Zielvideo", "Evidenzgebundenes Erklärvideo, etwa 6–10 Minuten"),
            ("Veröffentlichung", "Privater Upload; öffentliche Freigabe separat"),
        ],
        ("Baustein", "Live-Status"),
    )

    page_break(document)
    document.add_heading("1. Geplante und umgesetzte Vorgehensweise", level=1)
    document.add_paragraph(
        "Die Kanalbeschreibung wurde nicht als bloßer Schreibstil interpretiert, sondern in eine überprüfbare Produktionslogik übersetzt: verständlicher Einstieg, schrittweiser Wissensaufbau, klare Trennung von gesichertem Wissen, unterschiedlichen Einschätzungen und offenen Fragen sowie eine überprüfbare Quellenkette."
    )
    add_numbered(
        document,
        [
            "Ein wöchentliches Themenradar sucht konkrete, zeitgemäße Erklärfragen und unterdrückt ungeeignete Blickwinkel, Werbung, reine Meinung und Dubletten.",
            "Ein Mensch prüft Thema, Publikumsnutzen und Quellenpotenzial, bevor kostenintensive Recherche beginnt.",
            "Die KI zerlegt das Thema in mindestens sechs Erkläreinheiten. Quellen werden normalisiert und als unveränderliche Snapshots gesichert; Claims bleiben exakt an Fundstellen gebunden.",
            "Explanation Readiness prüft Wort- und Zeitkorridor, Claim-Dichte, Abdeckung aller wesentlichen Erkläreinheiten, unabhängige Quellen, Primärquellen und Gegenbelegsuche. Lücken lösen bis zu drei gezielte Evidenzrunden aus.",
            "Der Script Writer erzeugt ein deutsches, schrittweises Erklärskript ausschließlich aus freigegebenen Claims und Evidence-IDs.",
            "Ein separat zugewiesenes Verifier-Modell kontrolliert Belegtreue, Kausalität, Unsicherheit, Gegenpositionen und Verständlichkeit.",
            "Das Storyboard übersetzt jeden Erklärschritt in quellengebundene, zugängliche Visuals; synthetische Darstellungen realer Ereignisse dürfen nicht wie Dokumentation wirken.",
            "Nach Skript, Storyboard und exaktem Render greifen eigene menschliche Gates. Der Upload erfolgt privat, die Veröffentlichung separat.",
        ],
    )
    add_figure(
        document,
        "01-automation-overview.png",
        "Abbildung 1 – Der live angezeigte 13-stufige FaktischSimpel-Workflow mit automatischen, assistierten und menschlichen Schritten.",
        width=6.25,
    )

    page_break(document)
    document.add_heading("2. Sicherheits- und Qualitätsgrenzen", level=1)
    document.add_paragraph(
        "Kanal-Prompts werden zusätzlich zu den globalen Prompt-Verträgen geladen. Sie dürfen deshalb Ton, Aufbau und Erklärstrategie verfeinern, aber keine technischen oder redaktionellen Schutzregeln ersetzen."
    )
    add_table(
        document,
        [
            ("Freigegebene Claims", "Das Modell darf keine zusätzlichen Fakten erfinden oder aus allgemeinem Vorwissen ergänzen."),
            ("Exakte Belegbindung", "Jede Tatsachenbehauptung benötigt den identischen Narrationstext, korrekte Zeichenpositionen und eine freigegebene Evidence-ID."),
            ("Unabhängige Prüfung", "Writer und Verifier werden getrennt geroutet; der Verifier kann das Skript blockieren."),
            ("Deterministische Verträge", "JSON-Schema und nachgelagerte Validatoren bleiben maßgeblich."),
            ("Sechs Pflicht-Gates", "Opportunity, Dossier, Skript, Storyboard, exakter Render und Veröffentlichung."),
            ("Private-by-default", "Kein kanalbezogener Prompt kann automatische öffentliche Veröffentlichung aktivieren."),
        ],
        ("Grenze", "Erzwungenes Verhalten"),
    )
    add_note(
        document,
        "Prompt-Injection-Schutz",
        "Quellinhalt wird als nicht vertrauenswürdiges Beweismaterial behandelt. Anweisungen aus Webseiten oder Dokumenten dürfen keine System-, Freigabe- oder Ausgabegrenzen verändern.",
        accent=GOLD,
    )
    document.add_heading("Warum drei getrennte Prompts?", level=2)
    add_bullets(
        document,
        [
            "Der Writer optimiert Verständlichkeit und didaktische Reihenfolge.",
            "Der Verifier sucht aktiv nach unbelegten Aussagen, Übertreibung, falscher Kausalität und nur scheinbarer Ausgewogenheit.",
            "Der Storyboard-Generator prüft die visuelle Belegbarkeit und vermeidet Bilder, die mehr behaupten als die Quellen.",
        ],
    )
    add_figure(
        document,
        "02-writer-prompt.png",
        "Abbildung 2 – Der Writer-Prompt ist im Kanalprofil sichtbar und damit für Operatoren überprüfbar.",
        width=6.2,
    )

    page_break(document)
    document.add_heading("3. Themenradar konfigurieren", level=1)
    document.add_paragraph(
        "Das Profil wurde absichtlich deaktiviert erstellt, vollständig ergänzt, mit Test plan geprüft und erst danach aktiviert. Dadurch startet kein Zeitplan mit ungeprüften Suchparametern."
    )
    add_table(
        document,
        [
            ("Thema", subject["topic"]),
            ("Regionen", ", ".join(subject["regions"])),
            ("Frischefenster", f"{subject['freshness_policy']['lookback_days']} Tage"),
            ("Quellenminimum", "2 unabhängige Quellen, davon mindestens 1 Primärquelle"),
            ("Risikoprofil", "medium; supervised für Finance, Health und Politics"),
            ("Budget", "75.000 Tokens und 1.800 GPU-Sekunden pro Produktionsrahmen"),
            ("Format", "evidence_first_explainer, Ziel 480 Sekunden"),
        ],
        ("Feld", "Konfiguration"),
    )
    document.add_heading("Seed-Queries", level=2)
    add_bullets(document, subject["seed_queries"])
    document.add_heading("Ausgeschlossene Blickwinkel", level=2)
    add_bullets(document, subject["excluded_angles"])
    add_figure(
        document,
        "03-validated-search-plan.png",
        "Abbildung 3 – Der validierte Suchplan enthält Broad Discovery, Primärquellen-Suche, verwandte Konzepte und einen Falsifikationszweig.",
        width=6.2,
    )

    page_break(document)
    document.add_heading("4. Zeitplan und Live-Validierung", level=1)
    add_numbered(
        document,
        [
            "Nach erfolgreicher Planprüfung wurde das Subject aktiviert; damit entstand Profilversion 3.",
            "Apply schedule hat den Temporal-Zeitplan mit Cron 0 7 * * 1 und Europe/Berlin abgeglichen.",
            "Der nächste Lauf wurde für Montag, 27. Juli 2026, 07:00 Uhr Europe/Berlin eingeplant.",
            "Ein korrigierter manueller Live-Discovery-Lauf führte sieben private Metasuche-Abfragen nacheinander und mit Abstand aus. Aktuelle Suchen dürfen kontrolliert auf eine zeitlose Suche zurückfallen; Primär- und Gegenbelegzweige arbeiten grundsätzlich zeitlos.",
            "Der verifizierte Lauf sichtete 70 Rohresultate, verwarf 52 nicht erklärgeeignete Plattform- und Navigationstreffer, deduplizierte 18 geeignete Funde und erzeugte 11 neue, quellenverknüpfte Opportunities.",
        ],
    )
    add_figure(
        document,
        "04-subject-schedule-active.png",
        "Abbildung 4 – Das FaktischSimpel-Themenradar ist aktiviert und der Temporal-Zeitplan ist aktiv.",
        width=5.0,
    )
    add_figure(
        document,
        "06-live-discovery-completed.png",
        "Abbildung 5 – Der korrigierte Live-Lauf zeigt Orchestrierungsstatus und Suchquellen-Gesundheit getrennt; ArXiv, Bing und Bing News antworteten im Browser-Nachweis.",
        width=6.25,
    )
    add_note(
        document,
        "Korrektur des früheren Leer-Laufs",
        "Der frühere Null-Treffer-Lauf war kein fachlich leerer Themenmarkt. Brave war gedrosselt, DuckDuckGo und Startpage verlangten CAPTCHAs, SearXNG meldete diese Ausfälle aber zusammen mit HTTP 200. TubeFactory verwirft diese Diagnose nicht mehr: Null Treffer bei nicht erreichbaren Engines stoppt den Lauf als SOURCE_UNAVAILABLE mit Wiederholungsmöglichkeit. Teilabdeckung erscheint als degraded; nur belegte Quellenabdeckung erscheint healthy.",
    )

    page_break(document)
    document.add_heading("5. Bedienung im Redaktionsalltag", level=1)
    add_numbered(
        document,
        [
            "Im Channel workspace FaktischSimpel auswählen und in Channels & subjects prüfen, ob Workflow und Themenradar aktiv sind.",
            "In Research & evidence neue Findings vollständig öffnen. Nur ein Thema mit erkennbarem Publikumsnutzen und belastbarem Quellenpotenzial shortlist-en.",
            "Quellen akquirieren, Snapshots prüfen sowie exakte Claims, Gegenbelege und offene Fragen freigeben; anschließend die konkrete Dossier-Version genehmigen.",
            "Skript generieren. Der FaktischSimpel-Prompt wird aus der Kanal-Lineage automatisch an Writer und danach an den unabhängigen Verifier übergeben.",
            "Nur die verifizierte Skriptversion genehmigen. Danach Storyboard erzeugen, Quellenbindungen und Visuals prüfen und die exakte Version freigeben.",
            "Medien produzieren, blockierende QA-Findings lösen, exakten Render freigeben, privat hochladen und eine öffentliche Veröffentlichung separat autorisieren.",
        ],
    )
    document.add_heading("Suchbare Hilfe", level=2)
    document.add_paragraph(
        "Das globale Hilfe-Panel enthält jetzt die Aufgabe Produce a FaktischSimpel explainer. Die Suche nach FaktischSimpel reduziert die Liste auf diesen Ablauf; die sechs Schritte verlinken zurück in den Workflow."
    )
    add_figure(
        document,
        "05-task-help.png",
        "Abbildung 6 – Die FaktischSimpel-Prozedur ist im task-basierten Hilfe-Panel suchbar und vollständig aufklappbar.",
        width=5.1,
    )

    page_break(document)
    document.add_heading("6. Technische Umsetzung", level=1)
    add_table(
        document,
        [
            ("Validierter Blueprint", "config/channel-workflows/faktischsimpel.json"),
            ("Core Policy", "channel_workflow.py prüft Prompt-Längen, Aufgaben, Stufen und Pflicht-Gates."),
            ("API-Grenze", "Jede Channel-Version wird beim Schreiben gegen den Workflow-Vertrag validiert."),
            ("Lineage-Auflösung", "Dossier → Opportunity → Subject → Channel lädt die richtige Kanalversion für Skript und Verifier."),
            ("Storyboard-Lineage", "Script → Dossier → Opportunity → Subject → Channel lädt denselben Kanalworkflow für Szenen."),
            ("AI-Protokolltrennung", "OpenAI-kompatible Clients nutzen http://ai.b1.germering:11438/v1; native Ollama-Clients nutzen http://ai.b1.germering:11434."),
            ("Aktive Modellroute", "Alle sieben Editorial- und Research-Aufgaben verwenden gemma4-e4b:64k auf der OpenAI-kompatiblen Primärroute und demselben Modell über die native Ollama-Fallbackroute."),
            ("UI", "Das Kanalprofil zeigt Stufen, Status, Writer-, Verifier- und Storyboard-Prompt."),
            ("Suchquellen-Gesundheit", "SearXNG-Ergebnisse, antwortende Engines, Ausfälle, Fallbacks und Relevanzfilter-Zähler bleiben im Workflow-Ergebnis sichtbar."),
            ("Laststeuerung", "Sieben Strategien laufen sequenziell mit Abstand und begrenztem exponentiellem Retry statt als gleichzeitiger Anfrageburst."),
            ("Engine-Pools", "Allgemeine Recherche nutzt Bing, Qwant News und Bing News; Primär- und Gegenbelegsuche nutzt Bing, ArXiv und PubMed."),
            ("Explainer-Relevanz", "Für evidence_first_explainer werden Plattformnavigation und Funde ohne Erklär-, Studien-, Forschungs-, Wirkungs- oder Zusammenhangssignal vor dem Clustering verworfen."),
            ("Explanation Readiness", "Ein generisches Core-Gate verlangt für dieses Profil 810–1.350 Wörter, mindestens 13 nutzbare Tatsachenclaims, sechs Erkläreinheiten, unabhängige und primäre Quellen sowie eine Gegenbelegsuche."),
            ("Skalierbare Evidenzübergabe", "Für KI-Synthese und Dossierbildung werden deterministisch die 24 relevantesten unabhängigen Quellen ausgewählt; umfangreiche Semantic Chunks werden nicht über Temporal transportiert."),
            ("Neuheitspriorisierung", "Folgerunden bieten der KI zuvor ungenutzte wissenschaftliche Werke zuerst an; bereits verwendete Werke bleiben nur als Fallback verfügbar."),
            ("Autonome Kandidatenfolge", "Bleibt ein freigegebenes Thema nach drei Runden unzureichend, akquiriert und untersucht der Workflow selbstständig den nächsten freigegebenen Kandidaten. Harte Review-Gates werden dabei nicht gelockert."),
            ("Skript-Gates", "810–1.350 Wörter, 360–600 Sekunden und mindestens 13 unterschiedliche Tatsachenclaims werden nach der KI-Ausgabe deterministisch erzwungen."),
            ("Storyboard-Gates", "Konkrete On-Screen-Cues, narration-gebundene Visuals, wechselnde Visualtypen und höchstens 15 % Übergangsszenen werden deterministisch geprüft."),
            ("Regeneration", "Ein zuvor vorhandener Feldzugriffsfehler wurde korrigiert; Regeneration erhält Kanalprompt plus konkrete Änderungsanweisung."),
            ("Reproduzierbarer Build", "Veraltete interne Paketnamen nach der TubeFactory-Umbenennung wurden korrigiert."),
        ],
        ("Komponente", "Umsetzung"),
    )
    page_break(document)
    document.add_heading("Validierungsergebnisse", level=2)
    add_table(
        document,
        [
            ("Core", "58 Tests bestanden; einschließlich generischer Explanation-Readiness-Policy"),
            ("API", "112 Tests bestanden; einschließlich Versionierung, Readiness-Feldern und Skript-Sperre"),
            ("Research Worker", "66 Tests bestanden; einschließlich kumulativem Claim-Bestand, stabilen Erkläreinheiten, Quellen-Alias-Deduplizierung und Kandidatenübergabe"),
            ("Editorial Worker", "49 Tests bestanden; einschließlich stabiler Enrichment-Pläne, KI-Evidenzfeedback, Wiederholungssperre, priorisierter neuer Werke, begrenzter KI-Payloads, Activity-Heartbeat sowie Skript- und Storyboard-Hard-Gates"),
            ("Web", "TypeScript-Test und produktiver Next.js-Build bestanden"),
            ("Prompt-Grenzen", "Writer 2.389, Verifier 1.910, Storyboard 1.773 kompilierte Zeichen; alle unter 5.000"),
            ("Pflicht-Gates", "6 von 6 vorhanden"),
            ("Live-Datenbank", "Kanal v3 aktiv; Workflow 13 Stufen; Subject v3 aktiv; Cron und Zeitzone bestätigt"),
            ("Live Discovery", "70 Rohresultate; 18 geeignet; 52 irrelevant; 16 dedupliziert; 11 neue Opportunities; 3 antwortende Engines"),
            ("Live Dossier", "Lithium v11 explanation-ready: 14/13 nutzbare Claims, 7/6 Erkläreinheiten, 9 unabhängige und 6 Primärquellen, Gegenbelegsuche abgeschlossen"),
            ("Runtime", "API, Web, Research Worker und Editorial Worker healthy nach Deploy"),
        ],
        ("Prüfung", "Ergebnis"),
    )

    page_break(document)
    document.add_heading("7. Explanation Readiness und autonome Optimierung", level=1)
    document.add_paragraph(
        "Ein Dossier ist nicht mehr allein deshalb produktionsreif, weil wenige Claims formal belegt sind. "
        "Das neue Gate berechnet aus Zieldauer und Sprechtempo zunächst den benötigten Textumfang und daraus die minimale Zahl unterschiedlicher Fakten. "
        "Für FaktischSimpel ergibt das einen Korridor von 360–600 Sekunden beziehungsweise 810–1.350 Wörtern und mindestens 13 Tatsachenclaims."
    )
    add_table(
        document,
        [
            ("Erklärplan", "Mindestens 6 wesentliche Einheiten: Grundlage, Mechanismus, Evidenz, Grenzen, Folgen und offene Fragen"),
            ("Evidenz", "Mindestens 13 nutzbare atomare Claims bei 1,5 Claims je 100 Wörter"),
            ("Quellen", "Mindestens 2 unabhängige Quellen, davon 1 Primärquelle; Gegenbelegsuche muss ausgeführt sein"),
            ("Optimierung", "Bis zu 3 KI-geplante, lückenspezifische Such- und Syntheserunden über die LAN-KI"),
            ("Claim-Bestand", "Neue atomare Claims werden über alle Runden kumuliert; nahe Paraphrasen werden zusammengeführt und ihre exakten Evidence-IDs vereinigt"),
            ("Thesenrollen", "Maximal vier deterministisch geeignete Aussagen bleiben zentral; belastbare Definitionen, Mechanismen, Beispiele und Grenzen bleiben als unterstützende Claims erhalten"),
            ("Rückkopplung", "Die LAN-KI prüft den Zwischenstand quellengebunden; ihre Lückenhinweise steuern die nächste Suche, können aber kein deterministisches Gate überstimmen"),
            ("Abbruchschutz", "Meta-Claims über Dokumente, Quellenverfügbarkeit oder den Rechercheprozess zählen nicht als Fachbeleg"),
            ("Kandidatenwechsel", "Bei fortbestehenden Lücken automatische Übergabe an den nächsten bereits menschlich freigegebenen Kandidaten, maximal zwei Wechsel pro Reihe"),
        ],
        ("Gate", "Erzwungenes Verhalten"),
    )
    add_figure(
        document,
        "10-explanation-readiness-legacy-gate.png",
        "Abbildung 7 – Auch früher genehmigte Dossiers erscheinen als Needs evidence, wenn ihnen die neue vollständige Readiness-Auswertung fehlt.",
        width=6.2,
    )
    add_figure(
        document,
        "11-explanation-readiness-policy-fields.png",
        "Abbildung 8 – Dauer, Sprechtempo, maximale Evidenzrunden und Mindestzahl der Erkläreinheiten sind im Subject-Profil transparent konfigurierbar.",
        width=6.2,
    )
    add_figure(
        document,
        "12-explanation-readiness-result.png",
        "Abbildung 9 – Der reale Lithium-Test wurde nach drei LAN-KI-Runden korrekt blockiert: 4 von 13 Claims und fehlende wesentliche Erkläreinheiten reichen nicht für ein belastbares Video.",
        width=6.2,
    )
    add_figure(
        document,
        "13-autonomous-candidate-result.png",
        "Abbildung 10 – Nach dem autonomen Kandidatenwechsel akquirierte TubeFactory 22 Snapshots und deckte alle sechs Erkläreinheiten ab; mit 6 von 13 Claims blieb aber auch dieses Thema sicher vor der Skriptstufe blockiert.",
        width=6.2,
    )
    add_note(
        document,
        "Korrigierter generischer Ansatz",
        "Der Erklärplan wird zu Beginn einer Reihe einmal festgelegt und bleibt danach stabil. Jede Syntheserunde erhält den bisherigen Claim-Bestand, bearbeitet nur fehlende oder dünn belegte Einheiten und liefert zusätzliche atomare Aussagen. Bereits verwendete Evidence-IDs werden der Folgerunde nicht erneut angeboten. Quellen-Aliase werden über wissenschaftliche Werk-IDs zusammengeführt; zwei Auszüge desselben Werks zählen deshalb weiterhin nur als eine unabhängige Quelle. Schwächer belegte Details werden nicht mehr verworfen, sondern als nicht-zentrale Claims behalten, sofern mindestens ein exakter Direktbeleg vorhanden ist.",
    )
    add_note(
        document,
        "Warum das bessere Ergebnis zunächst ein Block ist",
        "Das frühere kurze Skript sah formal vollständig aus, hatte aber nur drei fachliche Claims. Das neue System verweigert diese Scheinsicherheit, benennt die konkreten Lücken und setzt die Recherche beziehungsweise Kandidatenauswahl automatisch fort.",
    )
    page_break(document)
    add_note(
        document,
        "Finaler autonomer Live-Nachweis",
        "Workflow live-research-dossier-corrected-claim-bank-feedback-v2-lithium-20260722a lief ohne manuelle Datenmanipulation bis DOSSIER_REVIEW. Dossier f67ac324-8ce5-4edc-bec8-dbd68aefd30f erreichte 14/13 nutzbare Claims, deckte alle sieben Erklärungseinheiten ab und belegte sie mit 9 unabhängigen sowie 6 Primärquellen. Die Gegenbelegsuche war abgeschlossen; das abschließende quellengebundene LAN-KI-Review bewertete 14 Claims und 9 Quellen ohne Abstention bei Konfidenz 1,0.",
    )
    add_figure(
        document,
        "14-corrected-claim-bank-result.png",
        "Abbildung 11 – Der vollständig korrigierte Lithium-Lauf ist explanation-ready: 14/13 Claims, sieben abgedeckte Erkläreinheiten, neun unabhängige und sechs Primärquellen sowie abgeschlossene Gegenbelegsuche.",
        width=6.2,
    )

    page_break(document)
    document.add_heading("8. Erhaltener Baseline-Lauf vor dem Readiness-Gate", level=1)
    document.add_paragraph(
        "Für die neue Reihe wurden die älteren FaktischSimpel-Findings recoverable archiviert. "
        "Das System entdeckte anschließend selbstständig ein neues Thema zur Lithiumgewinnung aus geothermischen Systemen. "
        "Die menschlichen Prüfschritte wurden für diesen Test bewusst übernommen; Skript- und Storyboardtexte wurden nicht manuell manipuliert."
    )
    add_table(
        document,
        [
            ("Ausgewähltes Thema", "Herausforderungen und Chancen für die Lithiumgewinnung aus geothermischen Systemen in Deutschland"),
            ("Evidenzdossier", "Historische Version 2; 3 atomare Claims; nach neuer Policy nicht mehr explanation-ready"),
            ("Skript", "Version 31 freigegeben; 8 Segmente; 410 Wörter; 182,2 Sekunden; 100 % Central-Claim-Coverage; 0 Fehler und 0 Warnungen"),
            ("Skript-Hash", "7f5e454c6330968b0cf6882306c88785dd56b9ce752d955e49dd1d17ee66846b"),
            ("Storyboard", "Version 8 freigegeben; 8 Szenen; 182,2 Sekunden; 7 visuelle Rollen mit Claim- und Quellenbindung"),
            ("Storyboard-Hash", "1a54a75c75303e2d058ec2a29ec65800d4551f9d1549ee158c59bab3ba14bb84"),
        ],
        ("Artefakt", "Live-Ergebnis"),
    )
    add_figure(
        document,
        "08-approved-script-storyboard-hashes.png",
        "Abbildung 12 – Die erhaltenen unveränderlichen Baseline-Versionen mit Abdeckung, Szenenzahl und Hash-Präfixen.",
        width=6.2,
    )
    document.add_heading("Automatisch behobene Qualitätsprobleme", level=2)
    add_bullets(
        document,
        [
            "HTTP-Antworten lokaler Modelle werden jetzt bis zum tatsächlichen Stream-Ende gelesen; ein einzelner Teil-Read kann valides JSON nicht mehr abschneiden.",
            "Writer-Ausgaben werden segmentweise erzeugt, von zwei AI-Prüfungen bewertet und in höchstens zwei begrenzten Korrekturpässen automatisch überarbeitet.",
            "Deterministische Gates blockieren weiterhin unbelegte Aussagen, unzulässige Quellenlücken und wiederholte Faktenblöcke, auch wenn ein Modell widersprüchliche Bewertungsfelder liefert.",
            "Wo kein geprüfter Gegenclaim existiert, erzeugt das System eine klar nicht-faktische Gegenbeleg- und Unsicherheitsstruktur statt eine Behauptung zu erfinden.",
            "Storyboards erhalten eine automatische visuelle Diversitätsplanung sowie source-bound Visual-Briefs; semantisch valide Alternativen werden lokal in den unveränderlichen SceneSpec-Vertrag eingebettet.",
            "Ein veralteter zweiter Editorial Worker wurde gestoppt, damit nur der aktuelle TubeFactory-Worker die Temporal-Queue verarbeitet.",
        ],
    )
    add_figure(
        document,
        "09-approved-script-v31.png",
        "Abbildung 13 – Das erhaltene Baseline-Skript v31 mit unabhängiger Prüfung, exaktem Hash und segmentweiser Evidenzzuordnung.",
        width=3.65,
    )
    page_break(document)
    add_figure(
        document,
        "07-approved-storyboard-v8.png",
        "Abbildung 14 – Das erhaltene Baseline-Storyboard v8: acht zeitlich abgestimmte Szenen mit unterschiedlichen Visualtypen und expliziten Claim-Bindungen.",
        width=6.0,
    )
    add_note(
        document,
        "Einordnung der Baseline",
        "Die unveränderlichen Artefakte bleiben als Audit-Baseline erhalten. Mit 410 Wörtern, 182 Sekunden und nur drei fachlichen Claims würden sie die neue 6–10-Minuten- und Explanation-Readiness-Policy nicht mehr passieren.",
    )

    page_break(document)
    document.add_heading("9. Grenzen und verantwortlicher Einsatz", level=1)
    add_bullets(
        document,
        [
            "Das Themenradar liefert Kandidaten, keine redaktionelle Wahrheit. Ein quellenbedingt leerer Lauf wird als Fehler sichtbar; ein nachweislich gesunder Leer-Lauf darf weiterhin keine künstlichen Themen erzeugen.",
            "Ein schwaches Dossier wird nicht durch Prompt-Rhetorik kaschiert: Es löst gezielte Evidenzrunden, einen autonomen Kandidatenwechsel oder einen sichtbaren Wartezustand aus.",
            "Gesundheit, Finanzen und Politik bleiben im supervised-Profil und benötigen besondere Sorgfalt.",
            "Analogien dienen nur dem Verständnis und niemals als Beleg.",
            "AI-generierte Visuals realer Ereignisse dürfen nicht fotorealistisch als dokumentarischer Beweis erscheinen und müssen als synthetisch gekennzeichnet werden.",
            "Vor jeder Veröffentlichung müssen exakte Skript-, Storyboard-, Render- und Metadatenversionen nachvollziehbar sein.",
        ],
    )
    document.add_heading("Routine-Checkliste", level=2)
    add_bullets(
        document,
        [
            "Workflow activity auf Needs attention prüfen.",
            "Themenradar-Karte auf schedule active kontrollieren.",
            "Neue Findings nach Nutzen, Evidenzpotenzial, Neuheit und Risiko vergleichen.",
            "Bei der Dossierprüfung Primärquelle, Unabhängigkeit, Gegenbelege und offene Fragen explizit kontrollieren.",
            "Explanation Readiness muss ready anzeigen; Needs evidence sperrt Skriptgenerierung unabhängig von früheren Freigaben.",
            "Nach jeder Regeneration die neue exakte Version erneut prüfen; frühere Freigaben gelten nicht automatisch weiter.",
            "Privaten Upload vollständig ansehen, bevor die Veröffentlichung separat autorisiert wird.",
        ],
    )

    page_break(document)
    document.add_heading("Anhang A – Exakte AI-Prompts", level=1)
    document.add_paragraph(
        "Die folgenden Prompts sind die kanalbezogenen Ergänzungen. Zur Laufzeit werden ihnen die unveränderlichen globalen Sicherheits-, Evidenz- und JSON-Verträge vorangestellt."
    )
    add_prompt(document, "A.1 Script Writer", prompts["script_writer"])
    add_prompt(document, "A.2 Independent Verifier", prompts["script_verifier"])
    add_prompt(document, "A.3 Storyboard", prompts["storyboard"])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
